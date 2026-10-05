"""Questionnaire schema + Word template round-trip for the requirements phase.

The .docx has one table per section with columns [ID | Question | Your answer]. On upload we match rows by
the ID column, so users can type freely in the answer cells. Any other document (a spec, notes, a PDF) is
read as free-text context for Echo instead.
"""
from __future__ import annotations

import io
from functools import lru_cache

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.agents.llm import prompts_path

SUGGEST = "__suggest__"  # answer sentinel: "not sure — let Echo propose something"
MAX_UPLOAD_CHARS = 120_000


@lru_cache
def questionnaire() -> dict:
    data = yaml.safe_load(prompts_path("intake_questionnaire.yaml").read_text(encoding="utf-8"))
    for s in data["sections"]:  # YAML turns bare Yes/No into booleans; options are always text
        for q in s["questions"]:
            if "options" in q:
                q["options"] = [str(o) for o in q["options"]]
    return data


def questions() -> dict[str, dict]:
    return {q["id"]: {**q, "section": s["title"]} for s in questionnaire()["sections"] for q in s["questions"]}


def completeness(answers: dict) -> dict:
    qs = questions()
    required = [q for q in qs.values() if q.get("required")]
    answered = [q for q in qs.values() if str(answers.get(q["id"], "")).strip()]
    req_done = [q for q in required if str(answers.get(q["id"], "")).strip()]
    return {"answered": len(answered), "total": len(qs), "required_done": len(req_done), "required_total": len(required),
            "missing_required": [q["id"] for q in required if q not in req_done]}


def answers_for_prompt(answers: dict) -> str:
    """Readable, grouped answers for the model (question labels included, unanswered marked)."""
    lines = []
    for s in questionnaire()["sections"]:
        lines.append(f"## {s['title']}")
        for q in s["questions"]:
            v = str(answers.get(q["id"], "")).strip()
            shown = "(user is not sure — please propose a sensible value)" if v == SUGGEST else (v or "(not answered)")
            lines.append(f"- [{q['id']}] {q['label']}{' *required*' if q.get('required') else ''}\n  → {shown}")
    known = questions()
    older = {k: v for k, v in answers.items() if k not in known and str(v).strip()}  # topics Echo asked before 10-05
    if older:
        lines.append("## Earlier answers (topics now asked by the specialists)")
        lines += [f"- [{k}] → {v}" for k, v in older.items()]
    return "\n".join(lines)


# ── Word template ──────────────────────────────────────────────────────────────
def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def build_template(project_name: str, answers: dict) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10.5)

    doc.add_heading(f"Requirement — {project_name}", level=0)
    intro = doc.add_paragraph()
    intro.add_run("How to use this template. ").bold = True
    intro.add_run("Fill in the “Your answer” column (add as much detail as you like — paste payloads, tables, notes). "
                  "Leave a cell empty if it doesn't apply, or write “not sure” and Echo will suggest something. "
                  "Do not change the ID column. Upload the file back in Orkestra → Requirement → Word template.")
    for s in questionnaire()["sections"]:
        doc.add_heading(s["title"], level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        for cell, text in zip(hdr, ["ID", "Question", "Your answer"]):
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.bold, run.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
            _shade(cell, "4B3FB8")
        for q in s["questions"]:
            row = table.add_row().cells
            id_run = row[0].paragraphs[0].add_run(q["id"])
            id_run.font.size, id_run.font.color.rgb = Pt(8), RGBColor(0x88, 0x88, 0x99)
            qp = row[1].paragraphs[0]
            qr = qp.add_run(q["label"] + (" *" if q.get("required") else ""))
            qr.bold = True
            hint = q.get("help") or (f"e.g. {q['example']}" if q.get("example") else "")
            if q.get("options"):
                hint = (hint + "\n" if hint else "") + "Options: " + " / ".join(q["options"])
            if hint:
                hr = row[1].add_paragraph().add_run(hint)
                hr.italic, hr.font.size, hr.font.color.rgb = True, Pt(8.5), RGBColor(0x66, 0x66, 0x77)
            v = str(answers.get(q["id"], "")).strip()
            row[2].text = "" if v == SUGGEST else v
        for r in table.rows:
            r.cells[0].width, r.cells[1].width, r.cells[2].width = Pt(70), Pt(190), Pt(230)
    doc.add_paragraph().add_run("* required").italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def parse_upload(filename: str, data: bytes) -> dict:
    """Returns {"answers": {...}} for our template, else {"text": "..."} as free context."""
    name = filename.lower()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        known, found = questions(), {}
        for table in doc.tables:
            for row in table.rows:
                cells = row.cells
                if len(cells) >= 3 and cells[0].text.strip() in known:
                    value = cells[2].text.strip()
                    if value:
                        found[cells[0].text.strip()] = SUGGEST if value.lower() in {"not sure", "n/s", "?"} else value
        if found:
            return {"answers": found}
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return {"text": "\n".join(parts)[:MAX_UPLOAD_CHARS]}
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return {"text": "\n".join((p.extract_text() or "") for p in reader.pages)[:MAX_UPLOAD_CHARS]}
    if name.endswith((".xlsx", ".xlsm")):  # mapping sheets: every sheet as rows of "a | b | c"
        try:
            return {"text": xlsx_text(data)[:MAX_UPLOAD_CHARS]}
        except Exception:  # noqa: BLE001 - not a readable workbook: kept as a file like any other
            pass
    text = data.decode("utf-8", errors="replace")
    if not text.strip() or (text.count("�") + text.count("\x00")) > max(8, len(text) // 100):
        # any kind of file is accepted (user, 10-04): an image, a zip, an old .xls… is kept for the crew, just not read
        return {"text": f"[{filename}: a {len(data):,}-byte file that isn't text (kept in inputs/{filename}; its content "
                        "can't be read here, so ask the user what it shows if it matters)]"}
    return {"text": text[:MAX_UPLOAD_CHARS]}


def xlsx_text(data: bytes) -> str:
    """The cells of an .xlsx (a zip of XML), sheet by sheet, without a spreadsheet library."""
    import re
    import zipfile
    from xml.etree import ElementTree

    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    z = zipfile.ZipFile(io.BytesIO(data))
    shared: list[str] = []
    if "xl/sharedStrings.xml" in z.namelist():
        for si in ElementTree.fromstring(z.read("xl/sharedStrings.xml")).findall("m:si", ns):
            shared.append("".join(t.text or "" for t in si.iter(f"{{{ns['m']}}}t")))
    out: list[str] = []
    sheets = sorted((n for n in z.namelist() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n)), key=lambda n: int(re.findall(r"\d+", n)[0]))
    for n in sheets:
        number = re.findall(r"\d+", n)[0]
        out.append(f"## Sheet {number}")
        for row in ElementTree.fromstring(z.read(n)).iter(f"{{{ns['m']}}}row"):
            cells = []
            for c in row.findall("m:c", ns):
                v = c.find("m:v", ns)
                if c.get("t") == "s" and v is not None and v.text and v.text.isdigit() and int(v.text) < len(shared):
                    cells.append(shared[int(v.text)])
                elif c.get("t") == "inlineStr":
                    cells.append("".join(t.text or "" for t in c.iter(f"{{{ns['m']}}}t")))
                else:
                    cells.append(v.text if v is not None and v.text else "")
            if any(x.strip() for x in cells):
                out.append(" | ".join(cells))
    return "\n".join(out)
